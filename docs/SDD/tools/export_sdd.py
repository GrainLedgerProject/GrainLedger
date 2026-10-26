"""Export the English SDD using the supplied Sources.docx typography and SDD structure.
Requires python-docx, Markdown and Pillow. Reference documents are read only.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import re,base64,html
from lxml import etree as ET
from docx import Document
from docx.shared import Inches,Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import markdown
ROOT=Path(__file__).resolve().parents[1];md=(ROOT/'GrainLedger_SDD.md').read_text()
doc=Document();ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
# Copy style definitions only, not HomeCycle text, images, numbering, headers or relationships.
with ZipFile(ROOT/'Sources.docx') as z:
 source=ET.fromstring(z.read('word/styles.xml'))
 for sid in ['Normal','Heading1','Heading2','Heading3','Heading4','Caption']:
  node=source.find(f'w:style[@w:styleId="{sid}"]',ns)
  old=doc.styles.element.find(f'w:style[@w:styleId="{sid}"]',ns)
  if node is not None:
   if old is not None:doc.styles.element.remove(old)
   doc.styles.element.append(deepcopy(node))
 default=source.find('w:docDefaults',ns);old=doc.styles.element.find('w:docDefaults',ns)
 if default is not None:
  if old is not None:doc.styles.element.remove(old)
  doc.styles.element.insert(0,deepcopy(default))
sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(25.4)
sec.header_distance=sec.footer_distance=Mm(12.5)
for p in sec.footer.paragraphs:
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
def plain(t):return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',t).replace('`','').replace('**','').replace('<br>','\n')
def paragraph(txt,style=None):
 p=doc.add_paragraph(style=style)
 for c in re.split(r'(\*\*.*?\*\*|`[^`]+`)',txt):
  r=p.add_run(plain(c))
  if c.startswith('**'):r.bold=True
  if c.startswith('`'):r.font.name='Consolas';r.font.size=Pt(9)
 return p
lines=md.splitlines();i=0;figures=0;tables=0
while i<len(lines):
 l=lines[i].strip()
 if not l:i+=1;continue
 if l.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   cells=[plain(c.strip()) for c in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
   i+=1
  t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Table Grid';tables+=1
  for j,v in enumerate(rows[0]):t.rows[0].cells[j].text=v
  for vals in rows[1:]:
   for c,v in zip(t.add_row().cells,vals):
    # line-break attributes/signatures rather than creating long inline cell paragraphs
    c.text=v.replace('; ',';\n')
  for idx,row in enumerate(t.rows):
   for c in row.cells:
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(4)
     for r in p.runs:r.font.size=Pt(9);r.bold=idx==0
   if idx==0:
    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
  continue
 m=re.match(r'!\[([^\]]+)\]\(([^)]+)\)',l)
 if m:
  pth=(ROOT/m[2]).with_suffix('.png');im=Image.open(pth);w,h=im.size
  # Inline figures and captions follow the sample: no added page/section per diagram.
  maxw,maxh=6.25,8.1
  fw=min(maxw,maxh*w/h)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
  p.add_run().add_picture(str(pth),width=Inches(fw));figures+=1
 elif re.match(r'^(Diagram|Table) \d+:',l):
  p=paragraph(l,'Caption');p.paragraph_format.keep_with_next=False
  for r in p.runs:r.italic=True;r.bold=False;r.font.size=Pt(10);r.font.color.rgb=RGBColor(0,0,0)
 elif l.startswith('#'):
  level=len(l)-len(l.lstrip('#'));txt=l[level:].strip()
  p=paragraph(txt,f'Heading {min(level,4)}')
  if level==2 and txt.startswith(('1. ','3. ')):p.paragraph_format.page_break_before=True
 elif l.startswith('- '):paragraph(l[2:],'List Bullet')
 else:
  p=paragraph(l)
  if not l.startswith('**'):p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 i+=1
doc.core_properties.title='GrainLedger — Software Design Description'
doc.core_properties.subject='Revision 0.3 — repository-informed full target design, Sources.docx presentation'
doc.core_properties.author='GrainLedger Project Team'
doc.save(ROOT/'GrainLedger_SDD.docx')
body=markdown.markdown(md,extensions=['tables'])
# Figure captions and table captions are visibly styled after the examples.
body=re.sub(r'<p>((?:Diagram|Table) \d+:[^<]+)</p>',r'<p class="caption">\1</p>',body)
for p in (ROOT/'diagrams'/'rendered').glob('*.svg'):
 body=body.replace('src="diagrams/rendered/'+p.name+'"','src="data:image/svg+xml;base64,'+base64.b64encode(p.read_bytes()).decode()+'"')
css='''body{max-width:1050px;margin:36px auto;padding:0 22px;font:16px/1.5 Calibri,Arial,sans-serif;color:#111;background:#fff}h1{color:#c00000;font-size:26px}h2{font-size:22px;margin-top:1.7em}h3{font-size:20px;margin-top:1.5em}h4{font-size:18px;margin-top:1.2em}table{border-collapse:collapse;width:100%;font-size:13px;margin:18px 0}th,td{border:1px solid #555;padding:8px;vertical-align:top;overflow-wrap:anywhere}th{font-weight:bold}img{display:block;max-width:100%;height:auto;margin:18px auto 5px;cursor:zoom-in}.caption{text-align:center;font-size:14px;font-style:italic;margin:6px 0 24px}code{background:#f3f3f3}dialog{border:1px solid #999;padding:14px;max-width:96vw;max-height:96vh}dialog img{max-width:none;margin:0}dialog button{position:sticky;top:0;padding:8px 16px}::backdrop{background:#333b}.hint{font-size:14px;color:#555}@media print{body{margin:0;max-width:none;font-size:11pt}h1,h2,h3,h4{break-after:avoid}h2{break-before:page}img{max-height:23cm;object-fit:contain;break-inside:avoid}.caption{break-before:avoid}table{font-size:9pt}.hint,dialog{display:none}}'''
script="const d=document.querySelector('dialog');document.querySelectorAll('main img').forEach(i=>i.onclick=()=>{d.querySelector('img').src=i.src;d.showModal()});d.querySelector('button').onclick=()=>d.close();d.onclick=e=>{if(e.target===d)d.close()};"
page='<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>GrainLedger SDD — Sources-aligned revision</title><style>'+css+'</style><p class="hint">Click a diagram to view its full resolution. All figure sources are supplied separately.</p><main>'+body+'</main><dialog><button>Close</button><img alt="Expanded design diagram"></dialog><script>'+script+'</script></html>'
(ROOT/'GrainLedger_SDD.html').write_text(page)
print('Exported sources-aligned DOCX/HTML:',figures,'figures and',tables,'tables; A4 portrait, source heading styles, inline figures and numbered captions.')
